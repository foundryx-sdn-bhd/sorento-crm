"""Dealer targets import, fifth real-file gate (lane SALES-TARGET-DEALER, crew 16 Oct, round 4).

Round 4 on the owner's 15 Oct prod backup: Mocha parses (135 of 135), every failing row stops at
name matching. Shapes, masked:
1. Mocha A/C II / A/C III rows: the dealer's ledgers are "[A/C I]" (level 1), "[CERAMIC]" and
   "[IBORN]" (no level). None answers to A/C II, yet the reason said "more than one customer has
   this name and account". The row stays a pick (owner ruling 16 Oct 09:12: nothing ambiguous is
   linked; A/C II = [CERAMIC] is an open owner question), with a reason that says what is true.
2. Two A/C I ledgers whose names differ only in punctuation or spacing ("TIP-TOP TRADING" and
   "TIP TOP TRADING"): the one whose name the file writes letter for letter is the exact match. When
   the file writes neither, it stays a pick.
3. THRESHOLD_REPEATED "A/C I 9% | A/C I 9%": the same threshold written twice with the same rate
   in one period is one tier, with a warning on the row.
"""
from __future__ import annotations

import io
from decimal import Decimal

from .test_sales_targets_s1 import BASE, _code, _uid, api, world  # noqa: F401  (fixtures)

D = Decimal


def _bytes(fill):
    from openpyxl import Workbook

    wb = Workbook()
    fill(wb.active)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _customer(db, company_id, name, level):
    from app.models.order import Customer

    row = Customer(id=_uid(), company_id=company_id, customer_code=_code(), customer_name=name, account_level=level)
    db.add(row)
    db.flush()
    return row


def _dry_run(client, db, name, account=None):
    def fill(ws):
        if account is None:
            ws["A1"], ws["B1"] = "DEALER NAME", "Q1"
            ws["A2"], ws["B2"] = name, "100,000 - 1%"
        else:
            ws["A1"], ws["B1"], ws["C1"] = "DEALER NAME", "ACCOUNT", "Q1"
            ws["A2"], ws["B2"], ws["C2"] = name, account, "100,000 - 1%"

    files = {"file": ("t.xlsx", _bytes(fill), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    db.autoflush = False
    try:
        return client.post(f"{BASE}/import", params={"dry_run": "true", "year": 2026}, files=files).json()["rows"][0]
    finally:
        db.autoflush = True


def _mocha_trio(db, company_id):
    return [
        _customer(db, company_id, "MIKE TILES SDN BHD [A/C I]", 1),
        _customer(db, company_id, "MIKE TILES SDN BHD [CERAMIC]", None),
        _customer(db, company_id, "MIKE TILES SDN BHD [IBORN]", None),
    ]


def test_an_account_no_ledger_of_the_name_has_is_not_called_several_matches(api):
    client, db, company_id = api
    trio = _mocha_trio(db, company_id)
    row = _dry_run(client, db, "MIKE TILES SDN BHD", "A/C II")
    assert row["status"] == "needs_pick"
    assert "more than one customer" not in row["reason"]
    assert row["reason"].startswith('no ledger of this name has the account "A/C II"'), row["reason"]
    # The dealer's own ledgers are the suggestions, so the pick is one click.
    assert {s["name"] for s in row["suggestions"]} == {c.customer_name for c in trio}


def test_a_suffix_account_no_ledger_has_reads_the_same(api):
    client, db, company_id = api
    _mocha_trio(db, company_id)
    row = _dry_run(client, db, "MIKE TILES SDN BHD A/C III")
    assert row["status"] == "needs_pick"
    assert row["reason"].startswith('no ledger of this name has the account "A/C III"'), row["reason"]


def test_the_file_s_account_still_auto_maps_beside_the_marker_ledgers(api):
    client, db, company_id = api
    one, _, _ = _mocha_trio(db, company_id)
    row = _dry_run(client, db, "MIKE TILES SDN BHD", "A/C I")
    assert row["match"] == "auto" and row["customer"] == one.customer_name


def test_a_name_no_ledger_has_keeps_its_own_reason(api):
    client, db, company_id = api
    _mocha_trio(db, company_id)
    row = _dry_run(client, db, "OSCAR STONE SDN BHD", "A/C I")
    assert row["reason"].startswith("no customer has this name and account"), row["reason"]


def test_the_ledger_the_file_writes_letter_for_letter_wins_a_punctuation_tie(api):
    client, db, company_id = api
    _customer(db, company_id, "TIP TOP TRADING SDN. BHD. [A/C I]", 1)
    hyphen = _customer(db, company_id, "TIP-TOP TRADING SDN BHD [A/C I]", 1)
    row = _dry_run(client, db, "Tip-Top Trading Sdn Bhd A/C I")
    assert row["match"] == "auto" and row["customer"] == hyphen.customer_name


def test_a_spacing_tie_breaks_the_same_way(api):
    client, db, company_id = api
    spaced = _customer(db, company_id, "HOME & DECO SDN BHD", 1)
    _customer(db, company_id, "HOME&DECO SDN BHD", 1)
    row = _dry_run(client, db, "home  &  deco sdn bhd", "A/C I")
    assert row["match"] == "auto" and row["customer"] == spaced.customer_name


def test_a_punctuation_tie_the_file_writes_neither_way_stays_a_pick(api):
    client, db, company_id = api
    _customer(db, company_id, "TIP TOP TRADING SDN. BHD. [A/C I]", 1)
    _customer(db, company_id, "TIP-TOP TRADING SDN BHD [A/C I]", 1)
    row = _dry_run(client, db, "TIP / TOP TRADING SDN BHD A/C I")
    assert row["status"] == "needs_pick"
    assert row["reason"].startswith("more than one customer has this name and account")
    assert len(row["suggestions"]) == 2


def _repeat(second_rate):
    from app.services.sales.dealer_target_import import parse_workbook

    def fill(ws):
        ws["A1"], ws["B1"] = "DEALER NAME", "Q1"
        ws.merge_cells("B1:E1")
        for col, text in zip("BCDE", ("Tier 1", "Tier 1 Incentive", "Tier 2", "Tier 2 Incentive")):
            ws[f"{col}2"] = text
        ws["A3"], ws["B3"], ws["C3"], ws["D3"], ws["E3"] = "X A/C I", 300000, "A/C I- 9%", 300000, second_rate

    return parse_workbook(_bytes(fill), year=2026)


def test_a_threshold_written_twice_with_the_same_rate_is_one_tier_with_a_warning():
    parsed = _repeat("A/C I- 9%")
    row = parsed.rows[0]
    assert row.status == "mapped", row.reason
    (q,) = [s for s in parsed.schemes if s.kind == "quarterly"]
    assert [(t.threshold, t.rate, dict(t.account_rates or {})) for t in q.tiers if not t.is_additional] == [
        (D("300000"), None, {next(iter(q.tiers[0].account_rates)): D("9")})
    ]
    assert any("written twice" in w for w in row.warnings), row.warnings


def test_a_threshold_written_twice_with_different_rates_still_fails():
    row = _repeat("A/C I- 7%").rows[0]
    assert row.status == "not_representable"
    assert "repeats in one period" in row.reason
