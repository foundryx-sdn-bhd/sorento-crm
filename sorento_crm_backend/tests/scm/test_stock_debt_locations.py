"""Stock Debt Location is the real warehouse (STOCK-DEBT-SAVED-VIEWS round 3, AC-R1..AC-R6).

Written BEFORE the implementation. `location` / `exclude_location` are repeatable warehouse
ids on the list and cell routes, lists on the export body; the list envelope carries a
`locations` facet. Every assertion is on rows seeded here (ZZT-prefixed): the shared DB holds
real warehouses, so no exact facet list and no global count is asserted.
"""
from __future__ import annotations

from fastapi.testclient import TestClient

from app.services.scm.supply_assignment import month_key
from tests.scm.conftest import requires_pg
from tests.scm.test_stock_debt_routes import (
    BASE,
    TODAY,
    _client,
    _demand,
    _months_ahead,
    _product,
    _row_of,
    _stock,
    _u,
    _warehouse,
)

pytestmark = requires_pg

NOW = month_key(TODAY)
DUE = month_key(_months_ahead(1))


def _two_bins(db):
    """A and B are bins of the SAME group (-IB): 40 demanded at A next month, 100 on hand
    at B. Whole book: B covers A. Only A: A owes its 40."""
    marker = f"ZZTSD{_u()[:6]}".upper()
    a = _warehouse(db, f"ZZTLA{_u()[:4]}-IB")
    b = _warehouse(db, f"ZZTLB{_u()[:4]}-IB")
    product = _product(db, f"{marker}-A")
    _stock(db, product, b, 100)
    _demand(db, product, a, qty=40, required_date=_months_ahead(1), so_number=f"{marker}-SO1")
    db.flush()
    return marker, product, a, b


def _get(c, marker, *pairs, **extra):
    params = [("query", marker), ("only_debt", "false"), *pairs, *extra.items()]
    resp = c.get(BASE, params=params)
    assert resp.status_code == 200, resp.text
    return resp.json()


def _balances(body, product):
    row = _row_of(body, product.product_code)
    return {m["key"]: m["balance"] for m in row["months"]}


def _balances_or_none(body, product):
    rows = [r for r in body["data"] if r["product_code"] == product.product_code]
    if not rows:
        return None
    return {m["key"]: m["balance"] for m in rows[0]["months"]}


def _reserve(db, qty_product=None):
    from app.models.inventory import Warehouse
    from tests.scm.conftest import SORENTO_COMPANY_ID

    hq = Warehouse(
        id=_u(), warehouse_code=f"ZZTHQ{_u()[:6]}", warehouse_name="ZZT HQ", location="ZZT",
        is_active=True, fulfilment_planning=False, segment="project",
        company_id=SORENTO_COMPANY_ID,
    )
    db.add(hq)
    db.flush()
    hq.pool_warehouse_id = hq.id
    db.flush()
    return hq


def _site_pool(db, *bins):
    pool = _warehouse(db, f"ZZTPOOL{_u()[:5]}", planning=False)
    for b in bins:
        b.pool_warehouse_id = pool.id
    db.flush()
    return pool


# --------------------------------------------------------------------------- AC-R1


def test_only_one_warehouse_not_its_group(scm_app):
    """AC-R1: whole book B covers A; `location=A` spans A alone, so A's month reads -40."""
    app, db = _client(scm_app)
    marker, product, a, b = _two_bins(db)
    with TestClient(app) as c:
        whole = _balances(_get(c, marker), product)
        only_a = _balances(_get(c, marker, ("location", a.id)), product)
    assert whole.get(DUE, 0) >= 0
    assert only_a.get(DUE) == -40
    assert only_a.get(NOW, 0) == 0


def test_repeated_location_spans_both(scm_app):
    """AC-R1: `location=A&location=B` spans both, so A is covered again."""
    app, db = _client(scm_app)
    marker, product, a, b = _two_bins(db)
    with TestClient(app) as c:
        both = _balances(_get(c, marker, ("location", a.id), ("location", b.id)), product)
    assert both.get(DUE, 0) >= 0
    assert both[NOW] == 60  # 100 on hand less the 40 drawn by A


def test_exclude_location_drops_that_warehouse(scm_app):
    """AC-R1: `exclude_location=B` takes B's stock out, A's 40 is owed."""
    app, db = _client(scm_app)
    marker, product, a, b = _two_bins(db)
    with TestClient(app) as c:
        got = _balances(_get(c, marker, ("exclude_location", b.id)), product)
    assert got.get(DUE) == -40


# --------------------------------------------------------------------------- AC-R2


def test_a_location_outside_the_span_matches_nothing(scm_app):
    """AC-R2: a made-up id, and an active but unflagged bin's id, alone give an EMPTY span:
    the seeded product is absent and the total is 0, never the whole book."""
    app, db = _client(scm_app)
    marker, product, a, b = _two_bins(db)
    unflagged = _warehouse(db, f"ZZTLU{_u()[:4]}-IB", planning=False)
    db.flush()
    with TestClient(app) as c:
        for pick in (_u(), unflagged.id):
            body = _get(c, marker, ("location", pick))
            assert [r for r in body["data"] if r["product_code"] == product.product_code] == []
            assert body["pagination"]["total"] == 0


# --------------------------------------------------------------------------- AC-R3


def test_a_site_pool_is_a_location_like_any_other(scm_app):
    """AC-R3: Book All, Only [pool] shows the pool's product; Only [A] does not."""
    app, db = _client(scm_app)
    marker, product, a, b = _two_bins(db)
    pool = _site_pool(db, a)
    pool_product = _product(db, f"{marker}-P")
    _stock(db, pool_product, pool, 500)
    _demand(db, pool_product, pool, qty=20, required_date=_months_ahead(1),
            so_number=f"{marker}-POOL")
    db.flush()
    with TestClient(app) as c:
        only_pool = _balances_or_none(_get(c, marker, ("location", pool.id)), pool_product)
        only_a = _balances_or_none(_get(c, marker, ("location", a.id)), pool_product)
    assert only_pool is not None and only_pool[NOW] == 480
    assert only_a is None or only_a.get(NOW, 0) == 0


def test_retail_book_honours_location_over_the_pools(scm_app):
    """AC-R3: `book=retail` honours `location`: narrowing to another pool drops this one."""
    app, db = _client(scm_app)
    marker, product, a, b = _two_bins(db)
    pool = _site_pool(db, a)
    other = _site_pool(db, b)
    pool_product = _product(db, f"{marker}-P")
    _stock(db, pool_product, pool, 500)
    db.flush()
    with TestClient(app) as c:
        mine = _balances_or_none(
            _get(c, marker, ("location", pool.id), book="retail"), pool_product)
        elsewhere = _balances_or_none(
            _get(c, marker, ("location", other.id), book="retail"), pool_product)
    assert mine is not None and mine[NOW] == 500
    assert elsewhere is None or elsewhere.get(NOW, 0) == 0


def test_an_hq_project_reserve_is_a_location_like_any_other(scm_app):
    """AC-R3: the reserve's 500 covers a project demand at A only when in span."""
    app, db = _client(scm_app)
    marker = f"ZZTSD{_u()[:6]}".upper()
    a = _warehouse(db, f"ZZTLA{_u()[:4]}-IB")
    b = _warehouse(db, f"ZZTLB{_u()[:4]}-IB")
    hq = _reserve(db)
    product = _product(db, f"{marker}-A")
    _stock(db, product, hq, 500)
    _demand(db, product, a, qty=40, required_date=_months_ahead(1), so_number=f"{marker}-SO1")
    db.flush()
    with TestClient(app) as c:
        only_a = _balances(_get(c, marker, ("location", a.id)), product)
        with_hq = _balances(_get(c, marker, ("location", a.id), ("location", hq.id)), product)
        except_b = _balances(_get(c, marker, ("exclude_location", b.id)), product)
    assert only_a.get(DUE) == -40
    assert with_hq.get(DUE, 0) >= 0
    assert except_b.get(DUE, 0) >= 0


# --------------------------------------------------------------------------- cell


def test_cell_foots_with_the_list_under_location_and_exclude_location(scm_app):
    """AC-R1: the drill reads the same span as the board under both params."""
    app, db = _client(scm_app)
    marker, product, a, b = _two_bins(db)
    with TestClient(app) as c:
        for pairs in ([("location", a.id)], [("exclude_location", b.id)]):
            board = _balances(_get(c, marker, *pairs), product)
            cell = c.get(f"{BASE}/{product.id}/cell", params=[("month", DUE), *pairs])
            now_cell = c.get(f"{BASE}/{product.id}/cell", params=[("month", NOW), *pairs])
            assert cell.status_code == 200, cell.text
            assert board.get(DUE) == -40
            demand = cell.json()["demand"]
            assert len(demand) == 1 and demand[0]["short_qty"] == 40
            on_hand = [r for r in now_cell.json()["supply"] if r["kind"] != "po"]
            assert sum(r["qty"] for r in on_hand) == 0
        both = c.get(f"{BASE}/{product.id}/cell",
                     params=[("month", NOW), ("location", a.id), ("location", b.id)])
        on_hand = [r for r in both.json()["supply"] if r["kind"] != "po"]
        assert sum(r["qty"] for r in on_hand) == 100


# --------------------------------------------------------------------------- AC-R4


def _capture_enqueue(monkeypatch):
    from app.services import queue_service

    calls = []

    def fake(func, *args, **kwargs):
        calls.append(args)
        return type("J", (), {"id": "fake-job-id"})()

    monkeypatch.setattr(queue_service, "enqueue_job", fake)
    return calls


def test_export_schema_coerces_location_and_exclude_location():
    """AC-R4: list kept, bare string wrapped, null empty, for both fields."""
    from app.schemas.stock_debt import StockDebtExportIn

    for field in ("location", "exclude_location"):
        assert getattr(StockDebtExportIn(**{field: ["x", "y"]}), field) == ["x", "y"]
        assert getattr(StockDebtExportIn(**{field: "x"}), field) == ["x"]
        assert getattr(StockDebtExportIn(**{field: None}), field) == []
        assert getattr(StockDebtExportIn(), field) == []


def test_export_route_carries_location_lists_to_the_task(scm_app, monkeypatch):
    """AC-R4: the body's picks reach the worker as lists."""
    calls = _capture_enqueue(monkeypatch)
    app, db = _client(scm_app)
    with TestClient(app) as c:
        resp = c.post(f"{BASE}/export", json={
            "location": "abc", "exclude_location": ["d", "e"], "split": "none"})
    assert resp.status_code == 201, resp.text
    assert calls[0][2]["location"] == ["abc"]
    assert calls[0][2]["exclude_location"] == ["d", "e"]


def test_service_export_and_list_span_the_same_warehouses(scm_app):
    """AC-R4: workbook and board agree under `location` and under `exclude_location`."""
    from io import BytesIO

    from openpyxl import load_workbook

    from app.services.scm.stock_debt_service import StockDebtService

    app, db = _client(scm_app)
    marker, product, a, b = _two_bins(db)
    svc = StockDebtService(db)
    for kwargs in ({"location": [a.id]}, {"exclude_location": [b.id]}):
        listing = svc.list(query=marker, only_debt=False, limit=50, **kwargs)
        assert _balances(listing, product).get(DUE) == -40
        blob, _ct, _fn, counts = svc.export(
            query=marker, only_debt=False, split="none", **kwargs)
        assert counts["rows"] == 1, counts
        ws = load_workbook(BytesIO(blob))["Stock debt"]
        header = [cell.value for cell in ws[1]]
        assert ws[2][header.index("Total")].value == _row_of(
            listing, product.product_code)["total"]


# --------------------------------------------------------------------------- AC-R5


def _facet(body):
    return {e["id"]: e for e in body["locations"]}


def test_locations_facet_follows_the_book_and_ignores_picks(scm_app):
    """AC-R5: `{id, code, name}` per warehouse in the Book's span, by code, unchanged by a
    `location` pick; an unflagged bin is absent."""
    app, db = _client(scm_app)
    marker, product, a, b = _two_bins(db)
    unflagged = _warehouse(db, f"ZZTLU{_u()[:4]}-IB", planning=False)
    pool = _site_pool(db, a)
    db.flush()
    with TestClient(app) as c:
        all_ = _get(c, marker)
        project = _get(c, marker, book="project")
        retail = _get(c, marker, book="retail")
        picked = _get(c, marker, ("location", a.id))
    facet = _facet(all_)
    assert set(facet[a.id]) >= {"id", "code", "name"}
    assert facet[a.id]["code"] == a.warehouse_code
    assert b.id in facet and pool.id in facet and unflagged.id not in facet
    assert pool.id not in _facet(project) and a.id in _facet(project)
    assert pool.id in _facet(retail) and a.id not in _facet(retail)
    mine = [e["code"] for e in all_["locations"] if e["id"] in (a.id, b.id)]
    assert mine == sorted(mine)
    assert picked["locations"] == all_["locations"]


# --------------------------------------------------------------------------- AC-B parity


def test_exclude_group_is_gone_and_single_group_still_narrows(scm_app):
    """Parity with main: `exclude_group` is not a param any more (ignored), `group` stays a
    single string that narrows the span."""
    app, db = _client(scm_app)
    marker, product, a, b = _two_bins(db)
    with TestClient(app) as c:
        ignored = _balances(_get(c, marker, ("exclude_group", "IB")), product)
        grouped = _balances(_get(c, marker, group="IB"), product)
    assert ignored.get(DUE, 0) >= 0, "exclude_group must no longer narrow anything"
    assert grouped[NOW] == 60


# --------------------------------------------------------------------------- AC-L1


def test_the_workbook_header_names_the_retail_reference_columns(scm_app):
    """AC-L1: the three BRW columns read `BRW Retail ...` in the workbook header."""
    from io import BytesIO

    from openpyxl import load_workbook

    from app.services.scm.stock_debt_service import StockDebtService

    app, db = _client(scm_app)
    marker, *_ = _two_bins(db)
    blob, *_ = StockDebtService(db).export(
        query=marker, only_debt=False, split="none", show_aux=True,
    )
    header = [c.value for c in load_workbook(BytesIO(blob))["Stock debt"][1]]
    for label in ("BRW Retail On Hand Qty", "BRW Retail PO Qty", "BRW Retail Incoming Qty"):
        assert label in header, header
