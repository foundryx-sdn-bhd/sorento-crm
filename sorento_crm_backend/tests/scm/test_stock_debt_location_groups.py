"""Stock Debt Location is multi-valued (STOCK-DEBT-SAVED-VIEWS, AC-B1..AC-B5).

Written BEFORE the implementation: today `group` is a single string on the list/cell routes,
the export body and `StockDebtService`. The seeded shapes and helpers come from
`test_stock_debt_routes`; nothing is borrowed from existing rows.

A warehouse code's group is the text after its FIRST hyphen
(`sales_agent_service.group_of_warehouse_code`), so f"ZZTBRW{x}-BB" is group BB.
"""
from __future__ import annotations

from fastapi.testclient import TestClient

from app.services.scm.supply_assignment import month_key
from tests.scm.conftest import requires_pg
from tests.scm.test_stock_debt_routes import (
    BASE,
    TODAY,
    _client,
    _demand,
    _months_ahead,
    _product,
    _row_of,
    _stock,
    _u,
    _warehouse,
)

pytestmark = requires_pg


def _seed_three_groups(db):
    """100 on hand at -IB, a 40 demand at -BB next month, 7 on hand at -IR."""
    marker = f"ZZTSD{_u()[:6]}".upper()
    bb = _warehouse(db, f"ZZTBRW{_u()[:4]}-BB")
    ib = _warehouse(db, f"ZZTBRW{_u()[:4]}-IB")
    ir = _warehouse(db, f"ZZTBRW{_u()[:4]}-IR")
    product = _product(db, f"{marker}-A")
    _stock(db, product, ib, 100)
    _stock(db, product, ir, 7)
    _demand(
        db, product, bb, qty=40, required_date=_months_ahead(1),
        so_number=f"{marker}-SO1",
    )
    db.flush()
    return marker, product, bb, ib, ir


def _balances(body, product):
    row = _row_of(body, product.product_code)
    return {m["key"]: m["balance"] for m in row["months"]}


# --------------------------------------------------------------------------- AC-B1


def test_repeated_group_spans_both_groups_and_leaves_the_third_out(scm_app):
    """AC-B1: `?group=BB&group=IB` spans the bins of BOTH groups and the balance is
    recomputed over that span: the IB 100 is in (current month 100), the IR 7 is out (not
    107). A single value is unchanged: BB alone has no stock in span, so 0."""
    app, db = _client(scm_app)
    marker, product, *_ = _seed_three_groups(db)
    due = month_key(_months_ahead(1))
    now = month_key(TODAY)

    with TestClient(app) as c:
        both = c.get(
            BASE,
            params=[("query", marker), ("only_debt", "false"), ("group", "BB"),
                    ("group", "IB")],
        )
        single = c.get(
            BASE,
            params=[("query", marker), ("only_debt", "false"), ("group", "BB")],
        )
    assert both.status_code == 200, both.text
    assert single.status_code == 200, single.text

    spanned = _balances(both.json(), product)
    assert spanned[now] == 100, "IB stock in span, IR 7 out of it"
    assert spanned.get(due) == -40
    narrowed = _balances(single.json(), product)
    assert narrowed[now] == 0
    assert narrowed.get(due) == -40


def test_group_matching_is_case_insensitive(scm_app):
    """AC-B1: `group=bb` behaves like `BB` (guard: may already pass today)."""
    app, db = _client(scm_app)
    marker, product, *_ = _seed_three_groups(db)
    now = month_key(TODAY)

    with TestClient(app) as c:
        lower = c.get(
            BASE, params=[("query", marker), ("only_debt", "false"), ("group", "bb")]
        ).json()
        upper = c.get(
            BASE, params=[("query", marker), ("only_debt", "false"), ("group", "BB")]
        ).json()
        lower_many = c.get(
            BASE,
            params=[("query", marker), ("only_debt", "false"), ("group", "bb"),
                    ("group", "ib")],
        )

    assert _balances(lower, product) == _balances(upper, product)
    assert _balances(lower, product)[now] == 0
    assert lower_many.status_code == 200, lower_many.text
    assert _balances(lower_many.json(), product)[now] == 100


# --------------------------------------------------------------------------- AC-B2


def test_cell_with_repeated_group_foots_with_the_board(scm_app):
    """AC-B2: the drill under `group=BB&group=IB` shows the same line the board's cell does:
    short its 40, nothing assigned (R40: IB's stock does not cover a BB line)."""
    app, db = _client(scm_app)
    marker, product, *_ = _seed_three_groups(db)
    due = month_key(_months_ahead(1))
    groups = [("group", "BB"), ("group", "IB")]

    with TestClient(app) as c:
        board = c.get(
            BASE, params=[("query", marker), ("only_debt", "false"), *groups]
        ).json()
        cell = c.get(
            f"{BASE}/{product.id}/cell", params=[("month", due), *groups]
        )
        # On hand is listed in its HOME month (the current one) unless a line pins it.
        now_cell = c.get(
            f"{BASE}/{product.id}/cell",
            params=[("month", month_key(TODAY)), *groups],
        )
    assert cell.status_code == 200, cell.text
    assert now_cell.status_code == 200, now_cell.text
    body = cell.json()

    assert _balances(board, product).get(due) == -40
    assert len(body["demand"]) == 1
    assert body["demand"][0]["status"] == "short"
    assert body["demand"][0]["short_qty"] == 40
    assert body["demand"][0]["assigned_qty"] == 0
    # The IB bin is in the span, so its 100 is listed as supply; the IR 7 is not.
    on_hand = [row for row in now_cell.json()["supply"] if row["kind"] != "po"]
    assert sum(row["qty"] for row in on_hand) == 100


# --------------------------------------------------------------------------- AC-B3


def _enqueue_stub(monkeypatch):
    from app.services import queue_service

    monkeypatch.setattr(
        queue_service, "enqueue_job",
        lambda func, *a, **k: type("J", (), {"id": "fake-job-id"})(),
    )


def test_export_body_accepts_a_list_of_groups(scm_app, monkeypatch):
    """AC-B3: `group` in the export body may be a list (201, not 422)."""
    _enqueue_stub(monkeypatch)
    app, db = _client(scm_app)
    with TestClient(app) as c:
        resp = c.post(f"{BASE}/export", json={"group": ["BB", "IB"], "split": "none"})
    assert resp.status_code == 201, resp.text


def test_export_body_still_accepts_a_bare_string_group(scm_app, monkeypatch):
    """AC-B3: an old caller sending `group: "BB"` is still accepted (guard)."""
    _enqueue_stub(monkeypatch)
    app, db = _client(scm_app)
    with TestClient(app) as c:
        resp = c.post(f"{BASE}/export", json={"group": "BB", "split": "none"})
    assert resp.status_code == 201, resp.text


def test_service_export_and_list_span_the_same_groups(scm_app):
    """AC-B3: the workbook spans the bins the board does. `export(group=[...])` and
    `list(group=[...])` on the service give the same rows and the same total."""
    from io import BytesIO

    from openpyxl import load_workbook

    from app.services.scm.stock_debt_service import StockDebtService

    app, db = _client(scm_app)
    marker, product, *_ = _seed_three_groups(db)
    svc = StockDebtService(db)

    listing = svc.list(
        query=marker, only_debt=False, group=["BB", "IB"], limit=50,
    )
    assert _balances(listing, product)[month_key(TODAY)] == 100

    blob, _ct, _fn, counts = svc.export(
        query=marker, only_debt=False, group=["BB", "IB"], split="none",
    )
    assert counts["rows"] == 1, counts
    ws = load_workbook(BytesIO(blob))["Stock debt"]
    header = [cell.value for cell in ws[1]]
    data_row = [cell.value for cell in ws[2]]
    assert data_row[0] == product.product_code
    # Total column foots with the list row's own total under the same span.
    assert data_row[header.index("Total")] == _row_of(listing, product.product_code)["total"]


# --------------------------------------------------------------------------- AC-B4


def test_groups_facet_is_generic_and_needs_the_planning_flag(scm_app):
    """AC-B4: a flagged active `-FG` bin puts FG in `groups` with no code change; an
    UNflagged `-ED` bin does not put ED there. (The DB holds other data: membership only.)
    A -ED bin is created flagged=False and nothing else in this transaction is -ED."""
    app, db = _client(scm_app)
    marker = f"ZZTSD{_u()[:6]}".upper()
    _warehouse(db, f"ZZTBRW{_u()[:4]}-FG")
    _warehouse(db, f"ZZTBRW{_u()[:4]}-ED", planning=False)
    db.flush()

    with TestClient(app) as c:
        body = c.get(BASE, params={"query": marker}).json()

    assert "FG" in body["groups"], body["groups"]
    assert body["groups"] == sorted(body["groups"])
    from sqlalchemy import text

    flagged_ed = db.execute(
        text(
            "SELECT count(*) FROM warehouses WHERE is_active AND fulfilment_planning "
            "AND upper(split_part(warehouse_code, '-', 2)) = 'ED'"
        )
    ).scalar()
    if not flagged_ed:
        assert "ED" not in body["groups"], body["groups"]


# --------------------------------------------------------------------------- AC-B5


def test_retail_book_ignores_any_number_of_groups(scm_app):
    """AC-B5: `book=retail` with two groups equals `book=retail` with none."""
    app, db = _client(scm_app)
    marker, *_ = _seed_three_groups(db)
    base = [("query", marker), ("only_debt", "false"), ("book", "retail")]

    with TestClient(app) as c:
        none = c.get(BASE, params=base)
        many = c.get(BASE, params=[*base, ("group", "BB"), ("group", "IB")])
    assert none.status_code == 200, none.text
    assert many.status_code == 200, many.text
    assert many.json()["data"] == none.json()["data"]
    assert many.json()["pagination"] == none.json()["pagination"]
