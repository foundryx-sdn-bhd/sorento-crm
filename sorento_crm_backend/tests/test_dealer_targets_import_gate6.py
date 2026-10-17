"""Dealer targets import, sixth real-file gate (lane SALES-TARGET-DEALER, crew 17 Oct, round 5).

Round 5 on the owner's 15 Oct prod backup, masked:
1. The dry-run script never installed the company scope listeners, so `company_scope(...)`
   filtered nothing and the name index held BOTH companies' customers: every tie paired an SRT
   ledger with its MCH twin, and two rows mapped to the other company. The import now names
   its one acting company and reads only that company's customers, whatever the listeners do.
2. A file name with a bracketed branch ("<NAME> (<branch>)") auto-mapped to the one ledger of
   that name, which has no such branch. Not exact, so it waits on a pick (owner ruling 16 Oct).
3. Three rows of one dealer all writing A/C I made one target list one ledger three times: the
   raw "An account is listed twice." is now a row reason naming the rows.
4. Masked OTHER reasons kept their words but x'd the row numbers, which are layout.
"""
from __future__ import annotations

import io

import pytest

from .test_sales_targets_s1 import BASE, _code, _uid, api, world  # noqa: F401  (fixtures)


def _bytes(fill):
    from openpyxl import Workbook

    wb = Workbook()
    fill(wb.active)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _customer(db, company_id, name, level):
    from app.models.order import Customer

    row = Customer(id=_uid(), company_id=company_id, customer_code=_code(), customer_name=name, account_level=level)
    db.add(row)
    db.flush()
    return row


def _one_row(name, account="A/C I"):
    def fill(ws):
        ws["A1"], ws["B1"], ws["C1"] = "DEALER NAME", "ACCOUNT", "Q1"
        ws["A2"], ws["B2"], ws["C2"] = name, account, "100,000 - 1%"

    return _bytes(fill)


def _post(client, db, data):
    files = {"file": ("t.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    db.autoflush = False
    try:
        return client.post(f"{BASE}/import", params={"dry_run": "true", "year": 2026}, files=files).json()
    finally:
        db.autoflush = True


def _other_company(db):
    from app.models.company import Company

    other = Company(id=_uid(), name="ZZT Other Co", code=f"Z{_uid()[:6]}")
    db.add(other)
    db.flush()
    return other


def test_the_import_refuses_to_run_without_one_acting_company(api):
    from app.models.base import company_scope
    from app.services.error_handler import AppException
    from app.services.sales.dealer_target_import import run

    _client, db, _ = api
    with company_scope(db, None), pytest.raises(AppException) as err:
        run(db, _one_row("ALFA HOME"), year=2026, dry_run=True)
    assert err.value.status_code == 422
    assert err.value.code == "IMPORT_NEEDS_ONE_COMPANY"


def test_the_name_index_reads_only_the_acting_company_whatever_the_scope(api):
    """Guard against an unscoped session (the dry-run script had no listeners)."""
    from app.models.base import company_scope
    from app.services.sales.dealer_target_import import _CustomerIndex

    _client, db, company_id = api
    other = _other_company(db)
    own = _customer(db, company_id, "BRAVO HOME [A/C I]", 1)
    with company_scope(db, None):
        _customer(db, other.id, "BRAVO HOME - [A/C I]", 1)
        exact, suggestions, _capped, _why = _CustomerIndex(db, company_id).match("BRAVO HOME", "A/C I")
    assert exact is not None and exact.id == str(own.id)
    assert suggestions == []


def test_the_dry_run_script_installs_the_company_scope_listeners(monkeypatch, tmp_path):
    from app.services import company_scope as scope_module
    from scripts import dealer_targets_dry_run

    class Installed(Exception):
        pass

    def spy():
        raise Installed

    monkeypatch.setattr(scope_module, "register_company_scope_listeners", spy)
    path = tmp_path / "t.xlsx"
    path.write_bytes(_one_row("ALFA HOME"))
    with pytest.raises(Installed):
        dealer_targets_dry_run.main([str(path), "--company", "NOPE", "--year", "2026"])


def test_a_branch_no_ledger_has_waits_on_a_pick(api):
    client, db, company_id = api
    plain = _customer(db, company_id, "XRAY HOME SDN BHD [A/C I]", 1)
    row = _post(client, db, _one_row("XRAY HOME SDN BHD (KL)"))["rows"][0]
    assert row["status"] == "needs_pick", row
    assert row["reason"].startswith('no ledger of this name has the branch "KL"'), row["reason"]
    assert [s["name"] for s in row["suggestions"]] == [plain.customer_name]


def test_a_branch_the_ledger_has_still_auto_maps(api):
    client, db, company_id = api
    kl = _customer(db, company_id, "XRAY HOME SDN BHD (KL) [A/C I]", 1)
    row = _post(client, db, _one_row("XRAY HOME SDN BHD (KL)"))["rows"][0]
    assert row["match"] == "auto" and row["customer"] == kl.customer_name


def test_one_ledger_written_on_three_rows_of_a_target_names_the_rows(api):
    client, db, company_id = api
    _customer(db, company_id, "ZULU HOME SDN BHD [A/C I]", 1)

    def fill(ws):
        ws["A1"], ws["B1"], ws["C1"] = "DEALER NAME", "ACCOUNT", "Q1"
        ws["A2"] = "ZULU HOME SDN BHD"
        ws.merge_cells("A2:A4")
        ws["C2"] = "100,000 - 1%"
        ws.merge_cells("C2:C4")
        for r in (2, 3, 4):
            ws[f"B{r}"] = "A/C I"

    rows = _post(client, db, _bytes(fill))["rows"]
    for row in rows:
        assert row["status"] == "not_representable", row
        assert "rows 2, 3 and 4 are one CRM customer" in row["reason"], row["reason"]


def test_masked_reasons_keep_row_numbers():
    from scripts.dealer_targets_dry_run import _mask

    assert "row 46" in _mask("the same accounts and dates as row 46")
    assert "rows 2, 3 and 4" in _mask("rows 2, 3 and 4 are one CRM customer; a target lists each account once")
