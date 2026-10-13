"""QUOTATION-UX backend red tests (UAC AC-QLG06 to AC-QLG12, AC-QLG14).

Written BEFORE the implementation, from the plan's API contract:

- the photo rule: one "Product Photos" image and no chosen photo IS the photo; a drawing
  ("Technical Specifications") is never a photo candidate;
- `quotation_lines.drawing_attachment_id` (nullable FK, ON DELETE SET NULL);
- line write validation of picked photo / drawing ids (422, unchanged stored id not re-checked);
- the line response: `drawing_attachment_id`, `technical_drawing`, `candidates`;
- issue freezes the drawing (fill only), revise copies both ids;
- the issue PDF and Excel print the drawing in TECHNICAL SPEC, legacy text under the description,
  description line breaks kept;
- the "A custom item needs a description" wording.

Postgres only, via ``blank_session``. Every row carries the ``zzt-qlg`` marker and is seeded here.
"""
from __future__ import annotations

import json
import re
import uuid
from datetime import datetime, timedelta
from decimal import Decimal
from io import BytesIO

from sqlalchemy import text

from app.models.product import Product, ProductAttachment
from app.models.resources import Attachment, AttachmentType

from ._pg_fixture import blank_session
from .test_project_quotation_bulk_lines import (  # noqa: F401  (api is a pytest fixture)
    _bulk,
    _quotation,
    _stored,
    api,
)
from .test_project_quotation_images import (
    _issue,
    _link,
    _priced_line,
    _product,
    _setup,
)

MARKER = "zzt-qlg"
PHOTOS = "Product Photos"
DRAWINGS = "Technical Specifications"
BASE = "/api/v1/project-sales"


def _uid() -> str:
    return str(uuid.uuid4())


def _type(db, name: str) -> AttachmentType:
    row = db.query(AttachmentType).filter(AttachmentType.type_name == name).first()
    if row is None:
        row = AttachmentType(
            id=_uid(),
            code=f"zzt{_uid()[:8]}",
            type_name=name,
            allowed_extensions="jpg,jpeg,png,pdf",
        )
        db.add(row)
        db.flush()
    return row


_clock = [datetime(2026, 1, 1)]


def _file(db, company_id, kind, name, *, mime="image/jpeg", deleted=False, age=None):
    """An attachment of the given kind. ``age`` orders created_at deterministically (the model
    stamps created_at from uploaded_at on insert, so uploaded_at is what is set)."""
    _clock[0] += timedelta(minutes=1)
    row = Attachment(
        id=_uid(),
        company_id=company_id,
        original_filename=f"{MARKER}-{name}",
        stored_filename=f"{MARKER}-{name}",
        file_path=f"https://cdn.zzt.test/products/{MARKER}/{name}",
        thumbnail_path=f"https://cdn.zzt.test/products/{MARKER}/{name}.thumb.jpg",
        mime_type=mime,
        is_deleted=deleted,
        attachment_type_id=_type(db, kind).id,
        uploaded_at=age if age is not None else _clock[0],
    )
    db.add(row)
    db.flush()
    return row


def _photo(db, env, product, name, **kw):
    row = _file(db, env["company_id"], PHOTOS, name, **kw)
    _link(db, product, row)
    return row


def _drawing(db, env, product, name, **kw):
    row = _file(db, env["company_id"], DRAWINGS, name, **kw)
    _link(db, product, row)
    return row


def _scope(db, env):
    from app.services import project_quotation_document_service as qdocs

    document = qdocs.create_document(db, project=env["project"], actor_user_id=env["owner"])
    scope = qdocs.add_scope(
        db, document=document, scope_label=f"{MARKER} Townhouse", actor_user_id=env["owner"]
    )
    return document, scope


# ------------------------------------------------------------------ 1. migration / model


def test_the_line_table_has_a_nullable_drawing_column_that_sets_null_on_delete():
    from app.models.projects import ProjectQuotationLine

    assert hasattr(ProjectQuotationLine, "drawing_attachment_id")
    with blank_session() as db:
        column = db.execute(
            text(
                "select is_nullable, data_type from information_schema.columns "
                "where table_schema='projects' and table_name='quotation_lines' "
                "and column_name='drawing_attachment_id'"
            )
        ).first()
        assert column is not None, "projects.quotation_lines.drawing_attachment_id is missing"
        assert column[0] == "YES"
        assert column[1] == "uuid"
        fk = db.execute(
            text(
                "select c.confdeltype, cl.relname from pg_constraint c "
                "join pg_class t on t.oid = c.conrelid "
                "join pg_namespace n on n.oid = t.relnamespace "
                "join pg_class cl on cl.oid = c.confrelid "
                "join pg_attribute a on a.attrelid = t.oid and a.attnum = any(c.conkey) "
                "where c.contype='f' and n.nspname='projects' and t.relname='quotation_lines' "
                "and a.attname='drawing_attachment_id'"
            )
        ).first()
        assert fk is not None, "no foreign key on drawing_attachment_id"
        assert fk[1] == "attachments"
        assert fk[0] == "n", "ON DELETE must be SET NULL"


# ------------------------------------------------------------------ 2. photo rule


def test_one_product_photo_and_nothing_chosen_is_the_photo_in_every_reader():
    from app.services import product_image_service as images

    with blank_session() as db:
        env = _setup(db)
        product = _product(db, env)
        only = _photo(db, env, product, "only.jpg")

        one = images.for_product(db, product.id)
        assert one.state == images.CHOSEN
        assert one.attachment_id == only.id
        assert images.images_for(db, [product.id])[product.id].attachment_id == only.id
        assert images.chosen_attachment_id(db, product.id) == only.id


def test_two_photos_and_nothing_chosen_stays_not_chosen():
    from app.services import product_image_service as images

    with blank_session() as db:
        env = _setup(db)
        product = _product(db, env)
        _photo(db, env, product, "a.jpg")
        _photo(db, env, product, "b.jpg")

        found = images.for_product(db, product.id)
        assert found.state == images.NOT_CHOSEN
        assert found.candidate_count == 2
        assert images.chosen_attachment_id(db, product.id) is None


def test_a_drawing_is_never_a_photo_candidate():
    from app.services import product_image_service as images

    with blank_session() as db:
        env = _setup(db)
        both = _product(db, env)
        photo = _photo(db, env, both, "wc.jpg")
        _drawing(db, env, both, "wc-drawing.jpg")
        only_drawing = _product(db, env)
        _drawing(db, env, only_drawing, "alone-drawing.jpg")

        picked = images.for_product(db, both.id)
        assert picked.state == images.CHOSEN
        assert picked.attachment_id == photo.id
        assert picked.candidate_count == 1

        none = images.for_product(db, only_drawing.id)
        assert none.state == images.NO_PHOTOS
        assert none.candidate_count == 0
        assert images.chosen_attachment_id(db, only_drawing.id) is None


def test_is_primary_still_wins_over_the_single_photo_default():
    from app.services import product_image_service as images

    with blank_session() as db:
        env = _setup(db)
        product = _product(db, env)
        _photo(db, env, product, "unmarked.jpg")
        marked = _file(db, env["company_id"], PHOTOS, "marked.jpg")
        link = _link(db, product, marked)
        images.choose(db, link)

        assert images.for_product(db, product.id).attachment_id == marked.id
        assert images.chosen_attachment_id(db, product.id) == marked.id


# ------------------------------------------------------------------ 3. write validation


def _line_body(product=None, **extra):
    body = {"unit_price": "100.00", "quantity": "1", **extra}
    if product is not None:
        body["product_id"] = product.id
    else:
        body["description_snapshot"] = f"{MARKER} bespoke"
    return body


def _fixture_for_write(api):
    client, db, company_id, _user_id, project = api
    env = {"company_id": company_id}
    from app.models.product import ProductCategory, UnitOfMeasure

    uom = UnitOfMeasure(id=_uid(), uom_code=f"ZZT{_uid()[:4]}", uom_name="Piece")
    category = ProductCategory(
        id=_uid(), category_code=f"ZZT-{_uid()[:8]}", category_name=f"{MARKER} Cat"
    )
    db.add_all([uom, category])
    db.flush()
    env["category"], env["uom"] = category, uom

    def make_product():
        row = Product(
            id=_uid(),
            product_code=f"ZZT-{_uid()[:8]}",
            product_name=f"{MARKER} Basin",
            description="Wall-hung basin",
            category_id=category.id,
            base_uom_id=uom.id,
            list_price=Decimal("100.00"),
        )
        db.add(row)
        db.flush()
        return row

    quotation = _quotation(client, project.id)
    return client, db, env, make_product, quotation["current_version_id"]


def _refusal_text(response) -> str:
    return json.dumps(response.json())


def test_a_photo_and_a_drawing_of_the_lines_own_product_are_accepted(api):
    client, db, env, make_product, version_id = _fixture_for_write(api)
    product = make_product()
    photo = _photo(db, env, product, "p.jpg")
    drawing = _drawing(db, env, product, "d.jpg")
    db.commit()

    saved = _bulk(
        client,
        version_id,
        [
            _line_body(
                product,
                image_attachment_id=photo.id,
                drawing_attachment_id=drawing.id,
            )
        ],
    )
    assert saved.status_code == 200, saved.text
    row = saved.json()["data"][0]
    assert row["image_attachment_id"] == photo.id
    assert row["drawing_attachment_id"] == drawing.id


def test_a_photo_of_another_product_a_deleted_one_or_a_non_image_is_refused(api):
    client, db, env, make_product, version_id = _fixture_for_write(api)
    mine, other = make_product(), make_product()
    theirs = _photo(db, env, other, "theirs.jpg")
    gone = _photo(db, env, mine, "gone.jpg", deleted=True)
    pdf = _photo(db, env, mine, "sheet.pdf", mime="application/pdf")
    db.commit()

    for bad in (theirs.id, gone.id, pdf.id):
        refused = _bulk(client, version_id, [_line_body(mine, image_attachment_id=bad)])
        assert refused.status_code == 422, refused.text
        assert refused.json()["code"] == "quotation_line_image_invalid"


def test_a_photo_on_a_line_with_no_product_is_refused(api):
    client, db, env, make_product, version_id = _fixture_for_write(api)
    product = make_product()
    photo = _photo(db, env, product, "p.jpg")
    db.commit()

    refused = _bulk(client, version_id, [_line_body(None, image_attachment_id=photo.id)])
    assert refused.status_code == 422, refused.text
    assert refused.json()["code"] == "quotation_line_image_invalid"


def test_a_bad_drawing_is_refused_with_its_own_code(api):
    client, db, env, make_product, version_id = _fixture_for_write(api)
    mine, other = make_product(), make_product()
    theirs = _drawing(db, env, other, "theirs.jpg")
    gone = _drawing(db, env, mine, "gone.jpg", deleted=True)
    pdf = _drawing(db, env, mine, "sheet.pdf", mime="application/pdf")
    a_photo = _photo(db, env, mine, "photo.jpg")  # right product, wrong kind
    db.commit()

    for bad in (theirs.id, gone.id, pdf.id, a_photo.id):
        refused = _bulk(client, version_id, [_line_body(mine, drawing_attachment_id=bad)])
        assert refused.status_code == 422, refused.text
        assert refused.json()["code"] == "quotation_line_drawing_invalid"

    no_product = _bulk(
        client, version_id, [_line_body(None, drawing_attachment_id=theirs.id)]
    )
    assert no_product.status_code == 422, no_product.text
    assert no_product.json()["code"] == "quotation_line_drawing_invalid"


def test_an_unchanged_stored_id_is_not_rechecked(api):
    client, db, env, make_product, version_id = _fixture_for_write(api)
    product = make_product()
    photo = _photo(db, env, product, "p.jpg")
    drawing = _drawing(db, env, product, "d.jpg")
    db.commit()
    saved = _bulk(
        client,
        version_id,
        [_line_body(product, image_attachment_id=photo.id, drawing_attachment_id=drawing.id)],
    )
    assert saved.status_code == 200, saved.text
    line_id = saved.json()["data"][0]["id"]
    assert saved.json()["data"][0]["drawing_attachment_id"] == drawing.id

    # Both files are deleted afterwards; the issued-style record keeps its stamp.
    photo.is_deleted = True
    drawing.is_deleted = True
    db.commit()

    again = _bulk(
        client,
        version_id,
        [
            {
                "id": line_id,
                "image_attachment_id": photo.id,
                "drawing_attachment_id": drawing.id,
                "unit_price": "120.00",
            }
        ],
    )
    assert again.status_code == 200, again.text
    assert again.json()["data"][0]["drawing_attachment_id"] == drawing.id


# ------------------------------------------------------------------ 4. line response


def _lines(db, scope, owner, product, **payload):
    from app.services import project_quotation_service as quotes

    version = quotes.current_version(db, scope.id)
    line = quotes.upsert_line(
        db,
        version=version,
        actor_user_id=owner,
        payload={"unit_price": "10.00", "quantity": 1, "product_id": product.id, **payload}
        if product is not None
        else {"unit_price": "10.00", "quantity": 1, "description_snapshot": f"{MARKER} custom"},
    )
    return quotes.serialize_lines(db, [line])[0], line


def test_the_response_lists_photo_and_drawing_candidates_and_the_picked_drawing(api):
    client, db, env, make_product, version_id = _fixture_for_write(api)
    product = make_product()
    first = _drawing(db, env, product, "d1.jpg")
    second = _drawing(db, env, product, "d2.jpg")
    p1 = _photo(db, env, product, "p1.jpg")
    p2 = _photo(db, env, product, "p2.jpg")
    db.commit()

    saved = _bulk(client, version_id, [_line_body(product, drawing_attachment_id=second.id)])
    assert saved.status_code == 200, saved.text
    row = saved.json()["data"][0]
    assert row["drawing_attachment_id"] == second.id

    drawing = row["technical_drawing"]
    assert drawing["state"] == "chosen"
    assert drawing["attachment_id"] == second.id
    assert drawing["filename"] == f"{MARKER}-d2.jpg"
    assert drawing["candidate_count"] == 2
    assert [c["attachment_id"] for c in drawing["candidates"]] == [first.id, second.id]
    for candidate in drawing["candidates"]:
        assert set(candidate) >= {"attachment_id", "filename", "url", "preview_url"}

    photo = row["product_image"]
    assert photo["state"] == "not_chosen"
    assert photo["candidate_count"] == 2, "drawings are not photo candidates"
    assert {c["attachment_id"] for c in photo["candidates"]} == {p1.id, p2.id}


def test_an_unpicked_drawing_defaults_to_the_oldest_upload_and_none_says_so(api):
    client, db, env, make_product, version_id = _fixture_for_write(api)
    product, bare = make_product(), make_product()
    base = datetime(2025, 1, 1)
    _drawing(db, env, product, "newer.jpg", age=base + timedelta(days=5))
    oldest = _drawing(db, env, product, "older.jpg", age=base)
    db.commit()

    saved = _bulk(
        client, version_id, [_line_body(product), _line_body(bare), _line_body(None)]
    )
    assert saved.status_code == 200, saved.text
    with_two, without, custom = saved.json()["data"]

    assert with_two["drawing_attachment_id"] is None
    assert with_two["technical_drawing"]["state"] == "default"
    assert with_two["technical_drawing"]["attachment_id"] == oldest.id
    assert with_two["technical_drawing"]["candidate_count"] == 2
    assert without["technical_drawing"]["state"] == "no_drawings"
    assert without["technical_drawing"]["candidates"] == []
    assert custom["technical_drawing"]["state"] == "off_catalog"


def test_a_picked_photo_wins_over_the_products_live_answer(api):
    client, db, env, make_product, version_id = _fixture_for_write(api)
    product = make_product()
    a = _photo(db, env, product, "a.jpg")
    b = _photo(db, env, product, "b.jpg")
    db.commit()

    saved = _bulk(client, version_id, [_line_body(product, image_attachment_id=b.id)])
    assert saved.status_code == 200, saved.text
    cell = saved.json()["data"][0]["product_image"]
    assert cell["state"] == "chosen"
    assert cell["attachment_id"] == b.id
    assert a.id in {c["attachment_id"] for c in cell["candidates"]}


# ------------------------------------------------------------------ 5. freeze + revise


def test_issuing_fills_the_default_drawing_and_never_overwrites_a_picked_one():
    with blank_session() as db:
        env = _setup(db)
        defaulted, picked, custom = _product(db, env), _product(db, env), None
        base = datetime(2025, 1, 1)
        first = _drawing(db, env, defaulted, "first.jpg", age=base)
        _drawing(db, env, defaulted, "second.jpg", age=base + timedelta(days=1))
        _drawing(db, env, picked, "auto.jpg", age=base)
        chosen = _drawing(db, env, picked, "chosen.jpg", age=base + timedelta(days=1))

        document, scope = _scope(db, env)
        a = _priced_line(db, env, scope, defaulted)
        b = _priced_line(db, env, scope, picked, drawing_attachment_id=chosen.id)
        c = _priced_line(db, env, scope, custom)
        assert a.drawing_attachment_id is None

        _issue(db, document, env["owner"])
        for row in (a, b, c):
            db.refresh(row)
        assert a.drawing_attachment_id == first.id
        assert b.drawing_attachment_id == chosen.id
        assert c.drawing_attachment_id is None


def test_a_revision_copies_both_picked_ids():
    from app.services import project_quotation_service as quotes

    with blank_session() as db:
        env = _setup(db)
        product = _product(db, env)
        photo = _photo(db, env, product, "p.jpg")
        _photo(db, env, product, "p2.jpg")
        drawing = _drawing(db, env, product, "d.jpg")
        _document, scope = _scope(db, env)
        line = _priced_line(
            db,
            env,
            scope,
            product,
            image_attachment_id=photo.id,
            drawing_attachment_id=drawing.id,
        )
        assert line.drawing_attachment_id == drawing.id

        nxt = quotes.revise(db, quotation=scope, actor_user_id=env["owner"])
        copied = quotes.list_lines(db, nxt.id)
        assert [row.image_attachment_id for row in copied] == [photo.id]
        assert [row.drawing_attachment_id for row in copied] == [drawing.id]


# ------------------------------------------------------------------ 6. PDF and Excel


class _FakeBackend:
    def __init__(self, payload: bytes):
        self.payload = payload

    def download_file(self, key):
        return self.payload


def _jpeg() -> bytes:
    from PIL import Image

    buffer = BytesIO()
    Image.new("RGB", (80, 60), (200, 30, 30)).save(buffer, format="JPEG")
    return buffer.getvalue()


def _issued_with_drawing(db, monkeypatch, **payload):
    from app.services import product_image_service as images

    monkeypatch.setattr(images, "get_backend", lambda provider: _FakeBackend(_jpeg()))
    env = _setup(db)
    product = _product(db, env)
    drawing = _drawing(db, env, product, "d.jpg")
    document, scope = _scope(db, env)
    _priced_line(db, env, scope, product, drawing_attachment_id=drawing.id, **payload)
    return _issue(db, document, env["owner"])


def test_the_pdf_prints_the_drawing_in_technical_spec_and_legacy_text_under_the_description(
    monkeypatch,
):
    from app.services import project_quotation_pdf_service as qpdf

    with blank_session() as db:
        issue = _issued_with_drawing(
            db,
            monkeypatch,
            technical_spec=f"{MARKER} legacy spec words",
            description_snapshot=f"{MARKER} first row\nsecond row",
        )
        html = qpdf.build_issue_html(db, issue)

        spec = re.search(r'<td class="spec">(.*?)</td>', html, re.S)
        assert spec is not None
        assert '<img src="data:image/jpeg;base64,' in spec.group(1)
        assert "legacy spec words" not in spec.group(1)

        desc = re.search(r'<td class="desc">(.*?)</td>', html, re.S)
        assert desc is not None
        assert "legacy spec words" in desc.group(1)
        assert desc.group(1).index("first row") < desc.group(1).index("legacy spec words")


def test_the_pdf_keeps_description_line_breaks(monkeypatch):
    from app.services import project_quotation_pdf_service as qpdf

    with blank_session() as db:
        issue = _issued_with_drawing(
            db, monkeypatch, description_snapshot=f"{MARKER} row one\nrow two"
        )
        html = qpdf.build_issue_html(db, issue)
        desc = re.search(r'<td class="desc">(.*?)</td>', html, re.S).group(1)
        assert "row one" in desc and "row two" in desc
        assert "<br" in desc or "pre-line" in html or "pre-wrap" in html
        assert "row one row two" not in desc


def test_the_workbook_anchors_the_drawing_in_technical_spec_and_stacks_legacy_text(
    monkeypatch,
):
    from openpyxl.utils import get_column_letter

    from app.services import project_quotation_excel_service as qxlsx

    from .test_project_quotation_excel import _column_index, _data_rows

    with blank_session() as db:
        issue = _issued_with_drawing(
            db,
            monkeypatch,
            technical_spec=f"{MARKER} legacy spec words",
            description_snapshot=f"{MARKER} row one\nrow two",
        )
        sheet = qxlsx.build_issue_workbook(db, issue).worksheets[0]
        row = _data_rows(sheet)[0]
        spec_col = _column_index(sheet, "TECHNICAL SPEC")
        desc_col = _column_index(sheet, "DESCRIPTION")

        assert len(sheet._images) == 1
        assert sheet._images[0].anchor == f"{get_column_letter(spec_col)}{row}"
        assert sheet.cell(row=row, column=spec_col).value in (None, "")
        description = str(sheet.cell(row=row, column=desc_col).value)
        assert "row one\nrow two" in description
        assert "legacy spec words" in description


# ------------------------------------------------------------------ 7. wording


def test_a_custom_line_without_a_description_says_custom_item(api):
    client, _db, _company_id, _user_id, project = api
    version_id = _quotation(client, project.id)["current_version_id"]

    refused = _bulk(client, version_id, [{"unit_price": "10.00", "quantity": "1"}])
    assert refused.status_code == 422, refused.text
    body = _refusal_text(refused)
    assert refused.json()["code"] == "quotation_line_description_required"
    assert "A custom item needs a description" in body
    assert "off-catalog" not in body.lower()
