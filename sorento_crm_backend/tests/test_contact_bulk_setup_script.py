"""Lane CONTACT-BULK-SETUP: ``scripts/contact_bulk_setup.py`` (UAC
``documentation/plans/contacts/contact-bulk-setup-4oct-acceptance-criteria.md`` U1-U22).

RED tests, written before the script exists: the module import fails at collection
(ModuleNotFoundError), the right red. Fake data only. Postgres only (``blank_session``); every
company, warehouse, access type, agent, segment, contact and customer is seeded here.

Interface pinned (see the lane brief): ``RULES``; ``normalise_phone``; ``parse_workbook``;
``run`` (dry run by default; one savepoint per contact); ``write_csv``; ``links_review``;
``apply_links``. The per-contact audit row is written through the module attribute
``app.services.audit_service.record`` (event ``contact.bulk_setup``) so a test can make one
contact fail by patching it (U15); links audit as ``contact.customer_linked_bulk``.
"""
from __future__ import annotations

import csv
import uuid
from datetime import datetime
from types import SimpleNamespace

import openpyxl
import pytest

import scripts.contact_bulk_setup as cbs  # noqa: E402 - red until the script exists

from app.models.access import (
    AccessAgent,
    ContactAccessType,
    ContactAgentAccess,
    MarketSegment,
    RespondContact,
    RespondContactCustomer,
    StockVisibilityPolicy,
)
from app.models.audit import AuditLog
from app.models.base import set_company_scope
from app.models.company import Company, RespondContactCompany
from app.models.media import ContactMediaLimit
from app.models.order import Customer
from app.models.user import User
from app.services import audit_service
from app.services.company_scope import DEFAULT_COMPANY_ID

from tests._mc_lookup_seed import warehouse
from tests._pg_fixture import blank_session, unique_code

BO = "Business Owner /Director (Contain Cost Related Information)"
EMP = "Employee"
D1 = datetime(2026, 1, 1)
D2 = datetime(2026, 12, 31)
OTHER_FROM = datetime(2025, 5, 5)
OTHER_TO = datetime(2025, 6, 6)


def raw(n):  # as typed in the sheet
    return "0170000%04d" % n


def norm(n):
    return "+60170000%04d" % n


def dbphone(n):  # as stored on respond_contacts
    return "60170000%04d" % n


# ---------------------------------------------------------------- fixtures


@pytest.fixture
def db():
    with blank_session() as session:
        set_company_scope(session, frozenset({DEFAULT_COMPANY_ID}))
        yield session


@pytest.fixture
def w(db):
    """Two companies, their warehouses, access types, agents, segment and the injected RULES."""
    a = Company(id=str(uuid.uuid4()), name="ZZT Sorento Co", code="ZSRT")
    b = Company(id=str(uuid.uuid4()), name="ZZT Mocha Co", code="ZMCH")
    db.add_all([a, b])
    db.flush()
    set_company_scope(db, frozenset({DEFAULT_COMPANY_ID, a.id, b.id}))
    wh = {
        "ZBRW": warehouse(db, company_id=a.id, code="ZBRW"),
        "ZMWH": warehouse(db, company_id=a.id, code="ZMWH"),
        "ZHQ": warehouse(db, company_id=b.id, code="ZHQ"),
        "ZMOCHAWH": warehouse(db, company_id=b.id, code="ZMOCHAWH"),
    }
    types = {}
    for code in ("zzt_dealer", "zzt_mocha_dealer", "zzt_end_user", "zzt_extra"):
        t = ContactAccessType(code=code, name=f"ZZT {code}", is_active=True, sort_order=1)
        db.add(t)
        types[code] = t
    agents = [AccessAgent(code=f"zzt_ag_{i}", name=f"ZZT agent {i}") for i in range(3)]
    extra_agent = AccessAgent(code="zzt_ag_extra", name="ZZT agent extra")
    seg = MarketSegment(code="zzt_retail", name="ZZT retail", is_active=True, sort_order=1)
    seg_extra = MarketSegment(code="zzt_project", name="ZZT project", is_active=True, sort_order=2)
    db.add_all([*agents, extra_agent, seg, seg_extra])
    db.flush()
    rules = {
        "companies": {"Sorento Sdn Bhd": "ZSRT", "Mocha Sdn Bhd": "ZMCH"},
        "owner_types": {"ZSRT": "zzt_dealer", "ZMCH": "zzt_mocha_dealer"},
        "base_type": "zzt_end_user",
        "locations": {"ZSRT": ["ZBRW", "ZMWH"], "ZMCH": ["ZHQ", "ZMOCHAWH"]},
        "stock_mode": "availability",
        "segment": "zzt_retail",
        "media_modality": "image",
        "agents": [x.code for x in agents],
        "agent_valid_from": D1,
        "agent_valid_to": D2,
        "switches_on": ["chatbot_stock_allowed", "chatbot_eta_offset_applied"],
        "switches_off": ["escalation_allowed"],
    }
    actor = User(id=str(uuid.uuid4()), email=f"zzt-{uuid.uuid4().hex[:8]}@test.com", name="ZZT Actor")
    db.add(actor)
    db.flush()
    return SimpleNamespace(
        a=a, b=b, wh=wh, types=types, agents=agents, extra_agent=extra_agent, seg=seg,
        seg_extra=seg_extra, rules=rules, actor=actor.id,
    )


def _contact(db, n, **flags):
    row = RespondContact(id=str(uuid.uuid4()), phone_number=flags.pop("phone", dbphone(n)), name=f"ZZT Person {n}", **flags)
    db.add(row)
    db.flush()
    return row


def _customer(db, company, name, code=None):
    row = Customer(customer_code=code or unique_code("ZZT-D"), customer_name=name, company_id=company.id)
    db.add(row)
    db.flush()
    return row


HEADER = [None, "Dealer Registered Under", "Sales Person", "APPROVAL TO USE ", "Company Name"] + [
    h for i in range(1, 12) for h in (f"Name {i}", f"Contact Number {i}", f"Access Level {i}")
]


def _book(tmp_path, rows, name="working.xlsx"):
    """rows: dicts {approval, under, company, people: [(name, phone, level) | None, ...]}."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "working file"
    ws.append(HEADER)
    for r in rows:
        line = [None, r.get("under", "Sorento Sdn Bhd"), "ZZT Sales", r.get("approval", "YES "), r.get("company", "ZZT Dealer A Sdn Bhd")]
        people = list(r.get("people", []))
        for i in range(11):
            p = people[i] if i < len(people) else None
            line.extend(p if p else [None, None, None])
        ws.append(line)
    path = tmp_path / name
    wb.save(path)
    return path


def _row(people, **kw):
    return {"people": people, **kw}


def _state(db, contact_id):
    db.expire_all()
    c = db.get(RespondContact, contact_id)
    pol = db.query(StockVisibilityPolicy).filter(StockVisibilityPolicy.contact_id == contact_id).first()
    return {
        "types": sorted(t.code for t in c.access_types),
        "companies": sorted(
            r.company_id for r in db.query(RespondContactCompany).filter(RespondContactCompany.respond_contact_id == contact_id)
        ),
        "segments": sorted(s.code for s in c.market_segments),
        "stock": None if pol is None else (
            pol.mode, sorted(pol.warehouse_ids) if pol.warehouse_ids is not None else None,
            pol.excluded_warehouse_ids, pol.hide_zero_locations),
        "media": sorted(
            (m.modality, m.is_allowed, m.monthly_limit)
            for m in db.query(ContactMediaLimit).filter(ContactMediaLimit.contact_id == contact_id)
        ),
        "agents": sorted(
            (g.agent.code, g.is_allowed, g.valid_from, g.valid_to)
            for g in db.query(ContactAgentAccess).filter(ContactAgentAccess.respond_contact_id == contact_id)
        ),
        "flags": (c.chatbot_stock_allowed, c.chatbot_eta_offset_applied, c.escalation_allowed, c.notify_salesman, c.packing_list_allowed),
    }


def _audits(db, contact_id=None, event="contact.bulk_setup"):
    db.expire_all()
    q = db.query(AuditLog).filter(AuditLog.event == event)
    if contact_id:
        q = q.filter(AuditLog.entity_id == contact_id)
    return q.all()


def _facets(result, contact_id):
    return {c["facet"]: c for c in result["changes"] if c["contact_id"] == contact_id}


def _run(db, w, path, **kw):
    return cbs.run(db, path, rules=w.rules, **kw)


def _agent_rows(w):
    return sorted((x.code, True, D1, D2) for x in w.agents)


# ---------------------------------------------------------------- parse (U1-U5)


def test_u1_only_approved_rows_are_read(tmp_path, w):
    path = _book(tmp_path, [
        _row([("ZZT Yes", raw(1), BO)], approval="YES "),
        _row([("ZZT Lower", raw(2), BO)], approval=" yes"),
        _row([("ZZT Pending", raw(3), BO)], approval="PENDING"),
        _row([("ZZT Blank", raw(4), BO)], approval=None),
    ])
    entries, invalid = cbs.parse_workbook(path, rules=w.rules)
    assert sorted(e.phone for e in entries) == [norm(1), norm(2)]
    assert invalid == []


def test_u2_company_names_map_to_codes(tmp_path, w):
    path = _book(tmp_path, [
        _row([("ZZT A", raw(1), BO)], under="Sorento Sdn Bhd"),
        _row([("ZZT B", raw(2), BO)], under="Mocha Sdn Bhd"),
        _row([("ZZT C", raw(3), BO)], under="Sorento Sdn Bhd, Mocha Sdn Bhd"),
    ])
    entries, _ = cbs.parse_workbook(path, rules=w.rules)
    by = {e.phone: e for e in entries}
    assert by[norm(1)].companies == frozenset({"ZSRT"})
    assert by[norm(2)].companies == frozenset({"ZMCH"})
    assert by[norm(3)].companies == frozenset({"ZSRT", "ZMCH"})
    assert by[norm(1)].company_name == "ZZT Dealer A Sdn Bhd"
    assert by[norm(1)].name == "ZZT A" and by[norm(1)].level == "owner"
    assert by[norm(1)].row == 2 and by[norm(3)].row == 4


def test_u3_up_to_eleven_triples_with_gaps_skipped(tmp_path, w):
    people = [(f"ZZT P{i}", raw(i), BO if i % 2 else EMP) for i in range(1, 12)]
    people[2] = None  # an empty triple in the middle
    people[5] = None
    path = _book(tmp_path, [_row(people)])
    entries, invalid = cbs.parse_workbook(path, rules=w.rules)
    assert sorted(e.phone for e in entries) == sorted(norm(i) for i in (1, 2, 4, 5, 7, 8, 9, 10, 11))
    assert {e.phone: e.level for e in entries}[norm(11)] == "owner"
    assert {e.phone: e.level for e in entries}[norm(2)] == "employee"
    assert invalid == []


REAL_HEADER = [None, "Dealer Registered Under", "Sales Person", "APPROVAL TO USE ", "Company Name",
               "Name", "Contact Number", "Access Level"]
for _i in range(2, 6):
    REAL_HEADER += [f"Name {_i}", f"Contact Number {_i}", f"Access Level {_i}"]
for _i in range(6, 12):  # persons 6-11 reuse the "5" labels for phone and level
    REAL_HEADER += [f"Name {_i}", "Contact Number 5", "Access Level 5"]


def _real_book(tmp_path, slots_per_row):
    """slots_per_row: list of {slot(1..11): (name, phone, level)}; header exactly as the real sheet."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "working file"
    ws.append(REAL_HEADER)
    for slots in slots_per_row:
        line = [None, "Sorento Sdn Bhd", "ZZT Sales", "YES ", "ZZT Dealer A Sdn Bhd"]
        for i in range(1, 12):
            line.extend(slots.get(i, [None, None, None]))
        ws.append(line)
    path = tmp_path / "real.xlsx"
    wb.save(path)
    return path


def test_u3_real_header_first_person_unnumbered_and_persons_six_to_eleven_reuse_labels(tmp_path, w):
    path = _real_book(tmp_path, [
        {1: ("ZZT One", raw(1), BO), 5: ("ZZT Five", raw(5), EMP), 6: ("ZZT Six", raw(6), BO), 11: ("ZZT Eleven", raw(11), EMP)},
        {1: ("ZZT Solo", raw(7), BO)},
    ])
    entries, invalid = cbs.parse_workbook(path, rules=w.rules)
    got = {e.phone: (e.name, e.level, e.row) for e in entries}
    assert got == {
        norm(1): ("ZZT One", "owner", 2),
        norm(5): ("ZZT Five", "employee", 2),
        norm(6): ("ZZT Six", "owner", 2),
        norm(11): ("ZZT Eleven", "employee", 2),
        norm(7): ("ZZT Solo", "owner", 3),
    }
    assert invalid == []


@pytest.mark.parametrize(
    "value",
    ["0123456789", "60123456789", 60123456789, "+60 12-345 6789", "012 345 6789", "+60123456789"],
)
def test_u4_phone_formats_normalise(value):
    assert cbs.normalise_phone(value) == "+60123456789"


def test_u4_eight_and_nine_digit_tails_are_valid():
    assert cbs.normalise_phone("01111111111") == "+601111111111"  # +601 + 9 digits
    assert cbs.normalise_phone("0123456789") == "+60123456789"  # +601 + 8 digits


@pytest.mark.parametrize("value", ["12345", "0123456", "0312345678", "+6598765432", "", None, "abc", "012345678901"])
def test_u4_non_mobile_is_invalid(value):
    assert cbs.normalise_phone(value) is None


def test_u4_invalid_number_listed_never_matched(db, tmp_path, w):
    path = _book(tmp_path, [_row([("ZZT Good", raw(1), BO), ("ZZT Bad", "0312345678", BO)])])
    entries, invalid = cbs.parse_workbook(path, rules=w.rules)
    assert [e.phone for e in entries] == [norm(1)]
    (bad,) = invalid
    assert bad["row"] == 2 and bad["name"] == "ZZT Bad" and bad["reason"]
    assert bad["raw_phone_masked"] != "0312345678"
    assert set(bad) >= {"row", "name", "raw_phone_masked", "reason"}


def test_u5_same_phone_in_several_rows_is_the_union(db, tmp_path, w):
    c = _contact(db, 1)
    path = _book(tmp_path, [
        _row([("ZZT P1", raw(1), BO)], under="Sorento Sdn Bhd"),
        _row([("ZZT P1", raw(1), EMP)], under="Mocha Sdn Bhd"),
    ])
    res = _run(db, w, path, apply=True, actor_id=w.actor)
    s = _state(db, c.id)
    assert s["companies"] == sorted([w.a.id, w.b.id])
    # BO only for Sorento: Sorento dealer type yes, Mocha dealer type no (Employee there)
    assert s["types"] == ["zzt_dealer", "zzt_end_user"]
    assert not res["failed"]


def test_u5_same_phone_bo_and_employee_bo_wins(db, tmp_path, w):
    c = _contact(db, 1)
    path = _book(tmp_path, [_row([("ZZT P1", raw(1), EMP), ("ZZT P1", raw(1), BO)])])
    _run(db, w, path, apply=True, actor_id=w.actor)
    assert _state(db, c.id)["types"] == ["zzt_dealer", "zzt_end_user"]


# ---------------------------------------------------------------- desired state (U6, U7)


def test_u6_business_owner_gets_dealer_type_and_end_user(db, tmp_path, w):
    c = _contact(db, 1)
    _run(db, w, _book(tmp_path, [_row([("ZZT P1", raw(1), BO)])]), apply=True, actor_id=w.actor)
    assert _state(db, c.id)["types"] == ["zzt_dealer", "zzt_end_user"]


def test_u6_both_companies_owner_gets_both_dealer_types(db, tmp_path, w):
    c = _contact(db, 1)
    path = _book(tmp_path, [_row([("ZZT P1", raw(1), BO)], under="Sorento Sdn Bhd, Mocha Sdn Bhd")])
    _run(db, w, path, apply=True, actor_id=w.actor)
    assert _state(db, c.id)["types"] == ["zzt_dealer", "zzt_end_user", "zzt_mocha_dealer"]


def test_u6_employee_gets_end_user_only(db, tmp_path, w):
    c = _contact(db, 1)
    _run(db, w, _book(tmp_path, [_row([("ZZT P1", raw(1), EMP)])]), apply=True, actor_id=w.actor)
    assert _state(db, c.id)["types"] == ["zzt_end_user"]


def test_u7_full_desired_state_for_a_new_contact(db, tmp_path, w):
    c = _contact(db, 1, chatbot_stock_allowed=False, chatbot_eta_offset_applied=False, escalation_allowed=True)
    _run(db, w, _book(tmp_path, [_row([("ZZT P1", raw(1), BO)])]), apply=True, actor_id=w.actor)
    s = _state(db, c.id)
    assert s["companies"] == [w.a.id]
    assert s["stock"] == ("availability", sorted([w.wh["ZBRW"].id, w.wh["ZMWH"].id]), None, False)
    assert s["segments"] == ["zzt_retail"]
    assert s["media"] == [("image", True, None)]
    assert s["agents"] == _agent_rows(w)
    assert s["flags"][:3] == (True, True, False)


def test_u7_stock_locations_cover_every_desired_company(db, tmp_path, w):
    c = _contact(db, 1)
    path = _book(tmp_path, [_row([("ZZT P1", raw(1), EMP)], under="Sorento Sdn Bhd, Mocha Sdn Bhd")])
    _run(db, w, path, apply=True, actor_id=w.actor)
    assert _state(db, c.id)["stock"][1] == sorted(x.id for x in w.wh.values())


# ---------------------------------------------------------------- dry run (U8-U10)


def test_u8_dry_run_writes_nothing(db, tmp_path, w):
    c = _contact(db, 1, escalation_allowed=True)
    before = (_state(db, c.id), len(_audits(db)))
    res = _run(db, w, _book(tmp_path, [_row([("ZZT P1", raw(1), BO)])]))
    assert (_state(db, c.id), len(_audits(db))) == before
    assert res["changes"], "a dry run still reports what would change"
    assert all(ch["status"] in ("add", "set", "kept", "already") for ch in res["changes"])
    assert {ch["facet"] for ch in res["changes"]} >= {
        "access_types", "companies", "stock_mode", "stock_locations", "market_segments", "media",
        "agent_access", "chatbot_stock_allowed", "chatbot_eta_offset_applied", "escalation_allowed",
    }


def test_u9_csv_one_row_per_contact_and_facet_plus_missing(db, tmp_path, w):
    c = _contact(db, 1)
    path = _book(tmp_path, [_row([("ZZT P1", raw(1), BO), ("ZZT Ghost", raw(2), EMP), ("ZZT Bad", "0312345678", EMP)])])
    res = _run(db, w, path)
    out = tmp_path / "out"
    out.mkdir()
    changes_csv = out / "changes.csv"
    cbs.write_csv(res, changes_csv)
    with open(changes_csv, newline="") as fh:
        reader = csv.DictReader(fh)
        assert reader.fieldnames == ["phone", "contact", "excel_rows", "facet", "before", "after", "status"]
        rows = list(reader)
    mine = [r for r in rows if r["phone"] == norm(1)]
    assert len(mine) == len({r["facet"] for r in mine}) == len(_facets(res, c.id))
    assert all(r["contact"] == c.name and r["excel_rows"] == "2" for r in mine)
    others = [p for p in out.glob("*.csv") if p != changes_csv]
    assert len(others) == 1
    text = others[0].read_text()
    assert norm(2) in text and "ZZT Ghost" in text and "2" in text
    assert {m["reason"] for m in res["missing"]} >= {"missing", "invalid"}


def test_u10_contact_holding_everything_is_already_with_zero_adds(db, tmp_path, w):
    c = _contact(db, 1, escalation_allowed=False)
    c.access_types = [w.types["zzt_dealer"], w.types["zzt_end_user"]]
    c.market_segments = [w.seg]
    db.add(RespondContactCompany(respond_contact_id=c.id, company_id=w.a.id))
    db.add(StockVisibilityPolicy(
        id=str(uuid.uuid4()), contact_id=c.id, mode="availability", excluded_warehouse_ids=None,
        warehouse_ids=[w.wh["ZBRW"].id, w.wh["ZMWH"].id]))
    db.add(ContactMediaLimit(contact_id=c.id, modality="image", is_allowed=True))
    for ag in w.agents:
        db.add(ContactAgentAccess(respond_contact_id=c.id, respond_contact_phone=c.phone_number,
                                  agent_id=ag.id, is_allowed=True, valid_from=D1, valid_to=D2))
    db.flush()
    res = _run(db, w, _book(tmp_path, [_row([("ZZT P1", raw(1), BO)])]))
    mine = _facets(res, c.id)
    assert mine and {ch["status"] for ch in mine.values()} == {"already"}, mine


# ---------------------------------------------------------------- apply (U11-U18)


def test_u11_add_only_extras_are_never_removed(db, tmp_path, w):
    c = _contact(db, 1, notify_salesman=True, packing_list_allowed=True)
    other_co = Company(id=str(uuid.uuid4()), name="ZZT Other Co", code="ZOTH")
    db.add(other_co)
    db.flush()
    c.access_types = [w.types["zzt_extra"]]
    c.market_segments = [w.seg_extra]
    db.add(RespondContactCompany(respond_contact_id=c.id, company_id=other_co.id))
    db.add(ContactMediaLimit(contact_id=c.id, modality="voice", is_allowed=True, monthly_limit=7))
    db.add(ContactAgentAccess(respond_contact_id=c.id, respond_contact_phone=c.phone_number,
                              agent_id=w.extra_agent.id, is_allowed=True))
    db.flush()
    _run(db, w, _book(tmp_path, [_row([("ZZT P1", raw(1), BO)])]), apply=True, actor_id=w.actor)
    s = _state(db, c.id)
    assert s["types"] == ["zzt_dealer", "zzt_end_user", "zzt_extra"]
    assert s["segments"] == ["zzt_project", "zzt_retail"]
    assert s["companies"] == sorted([w.a.id, other_co.id])
    assert ("voice", True, 7) in s["media"] and ("image", True, None) in s["media"]
    assert ("zzt_ag_extra", True, None, None) in s["agents"]
    assert s["flags"][3:] == (True, True)


def test_u12_existing_agent_grant_with_other_dates_is_kept_and_reported(db, tmp_path, w):
    c = _contact(db, 1)
    db.add(ContactAgentAccess(respond_contact_id=c.id, respond_contact_phone=c.phone_number,
                              agent_id=w.agents[0].id, is_allowed=False, valid_from=OTHER_FROM, valid_to=OTHER_TO))
    db.flush()
    res = _run(db, w, _book(tmp_path, [_row([("ZZT P1", raw(1), BO)])]), apply=True, actor_id=w.actor)
    grants = {g[0]: g for g in _state(db, c.id)["agents"]}
    assert grants["zzt_ag_0"] == ("zzt_ag_0", False, OTHER_FROM, OTHER_TO)
    assert grants["zzt_ag_1"] == ("zzt_ag_1", True, D1, D2)
    assert _facets(res, c.id)["agent_access"]["status"] == "kept" or any(
        ch["facet"] == "agent_access" and ch["status"] == "kept" for ch in res["changes"] if ch["contact_id"] == c.id
    )


def test_u13_stock_override_is_set_not_merged(db, tmp_path, w):
    c = _contact(db, 1)
    gone = w.wh["ZHQ"].id
    db.add(StockVisibilityPolicy(
        id=str(uuid.uuid4()), contact_id=c.id, mode="compact", warehouse_ids=[gone],
        excluded_warehouse_ids=None, hide_zero_locations=True))
    db.flush()
    res = _run(db, w, _book(tmp_path, [_row([("ZZT P1", raw(1), BO)])]), apply=True, actor_id=w.actor)
    mode, include, excluded, hide = _state(db, c.id)["stock"]
    assert mode == "availability"
    assert include == sorted([w.wh["ZBRW"].id, w.wh["ZMWH"].id])  # the old Mocha location is gone
    assert excluded is None
    assert hide is True, "hide_zero_locations is kept on an existing row"
    ch = _facets(res, c.id)
    assert ch["stock_mode"]["status"] == "set"
    assert "compact" in str(ch["stock_mode"]["before"]) and "availability" in str(ch["stock_mode"]["after"])
    assert ch["stock_locations"]["status"] == "set"
    assert str(ch["stock_locations"]["before"]) != str(ch["stock_locations"]["after"])


def test_u13_existing_exclude_list_is_cleared(db, tmp_path, w):
    c = _contact(db, 1)
    db.add(StockVisibilityPolicy(
        id=str(uuid.uuid4()), contact_id=c.id, mode="availability", warehouse_ids=None,
        excluded_warehouse_ids=[w.wh["ZHQ"].id]))
    db.flush()
    _run(db, w, _book(tmp_path, [_row([("ZZT P1", raw(1), BO)])]), apply=True, actor_id=w.actor)
    assert _state(db, c.id)["stock"] == ("availability", sorted([w.wh["ZBRW"].id, w.wh["ZMWH"].id]), None, False)


def test_u13b_escalation_is_turned_off_and_reported(db, tmp_path, w):
    c = _contact(db, 1, escalation_allowed=True)
    res = _run(db, w, _book(tmp_path, [_row([("ZZT P1", raw(1), BO)])]), apply=True, actor_id=w.actor)
    assert _state(db, c.id)["flags"][2] is False
    ch = _facets(res, c.id)["escalation_allowed"]
    assert ch["status"] == "set" and ch["before"] != ch["after"]


def test_u14_image_media_row_not_allowed_is_set_true_limits_untouched(db, tmp_path, w):
    c = _contact(db, 1)
    db.add(ContactMediaLimit(contact_id=c.id, modality="image", is_allowed=False, monthly_limit=5))
    db.flush()
    _run(db, w, _book(tmp_path, [_row([("ZZT P1", raw(1), BO)])]), apply=True, actor_id=w.actor)
    assert _state(db, c.id)["media"] == [("image", True, 5)]


def test_u15_one_failing_contact_rolls_back_alone(db, tmp_path, w, monkeypatch):
    good1, bad, good2 = _contact(db, 1), _contact(db, 2), _contact(db, 3)
    before_bad = _state(db, bad.id)
    real = audit_service.record

    def boom(db_, **kw):
        if kw.get("entity_id") == bad.id:
            raise RuntimeError("simulated failure")
        return real(db_, **kw)

    monkeypatch.setattr(audit_service, "record", boom)
    path = _book(tmp_path, [_row([("ZZT P1", raw(1), BO), ("ZZT P2", raw(2), BO), ("ZZT P3", raw(3), BO)])])
    res = _run(db, w, path, apply=True, actor_id=w.actor)
    assert len(res["failed"]) == 1 and norm(2) in str(res["failed"][0])
    assert _state(db, bad.id) == before_bad, "the failing contact keeps its previous state on every facet"
    for ok in (good1, good2):
        assert _state(db, ok.id)["types"] == ["zzt_dealer", "zzt_end_user"]
        assert len(_audits(db, ok.id)) == 1
    assert _audits(db, bad.id) == []


def test_u16_one_audit_row_per_changed_contact(db, tmp_path, w):
    c1, c2 = _contact(db, 1), _contact(db, 2)
    path = _book(tmp_path, [_row([("ZZT P1", raw(1), BO), ("ZZT P2", raw(2), EMP)])])
    _run(db, w, path, apply=True, actor_id=w.actor)
    for c in (c1, c2):
        (row,) = _audits(db, c.id)
        assert row.entity_type == "respond_contacts" and row.action == "EVENT"
        assert row.old_values is not None and row.new_values
    assert "zzt_dealer" in str(_audits(db, c1.id)[0].new_values)


def test_u17_second_apply_is_idempotent(db, tmp_path, w):
    c = _contact(db, 1, escalation_allowed=True)
    path = _book(tmp_path, [_row([("ZZT P1", raw(1), BO)], under="Sorento Sdn Bhd, Mocha Sdn Bhd")])
    _run(db, w, path, apply=True, actor_id=w.actor)
    state, audits = _state(db, c.id), len(_audits(db))
    assert audits == 1
    res = _run(db, w, path, apply=True, actor_id=w.actor)
    assert _state(db, c.id) == state
    assert len(_audits(db)) == audits
    assert {ch["status"] for ch in res["changes"]} <= {"already", "kept"}


def test_u18_match_is_formatting_insensitive_and_never_creates(db, tmp_path, w):
    spaced = _contact(db, 1, phone="+60 17-0000 0001")
    before = db.query(RespondContact).count()
    path = _book(tmp_path, [_row([("ZZT P1", raw(1), BO), ("ZZT Nobody", raw(9), BO)])])
    res = _run(db, w, path, apply=True, actor_id=w.actor)
    assert _state(db, spaced.id)["types"] == ["zzt_dealer", "zzt_end_user"]
    assert db.query(RespondContact).count() == before
    assert [m["phone"] for m in res["missing"] if m["reason"] == "missing"] == [norm(9)]


# ---------------------------------------------------------------- links (U19-U22)


def _read_review(path):
    ws = openpyxl.load_workbook(path).active
    header = [c.value for c in ws[1]]
    rows = [dict(zip(header, [c.value for c in r])) for r in ws.iter_rows(min_row=2)]
    return header, [r for r in rows if r.get("excel_row") is not None]


def _link_world(db, tmp_path, w):
    c1, c2 = _contact(db, 1), _contact(db, 2)  # row 2: three people, one not in the DB
    d_exact = _customer(db, w.a, "ZZT Dealer A Sdn Bhd", "ZZT-D-EXACT")
    d_near = _customer(db, w.a, "ZZT Dealer A Trading", "ZZT-D-NEAR")
    for i in range(6):
        _customer(db, w.a, f"ZZT Unrelated Hardware {i}", f"ZZT-D-FAR{i}")
    m_only = _customer(db, w.b, "ZZT Dealer A Mocha Account", "ZZT-M-1")
    path = _book(tmp_path, [_row([("ZZT P1", raw(1), BO), ("ZZT P2", raw(2), EMP), ("ZZT Nobody", raw(9), EMP)])])
    return SimpleNamespace(c1=c1, c2=c2, exact=d_exact, near=d_near, m_only=m_only, path=path)


def _tick(review, reviewed, codes, mark="Y"):
    wb = openpyxl.load_workbook(review)
    ws = wb.active
    header = [c.value for c in ws[1]]
    ci, li = header.index("debtor_code"), header.index("link")
    for r in ws.iter_rows(min_row=2):
        if r[ci].value in codes:
            r[li].value = mark[r[ci].value] if isinstance(mark, dict) else mark
    wb.save(reviewed)
    return reviewed


def _links(db):
    db.expire_all()
    return db.query(RespondContactCustomer).all()


def test_u19_review_sheet_lists_up_to_five_ranked_candidates_per_company(db, tmp_path, w):
    lw = _link_world(db, tmp_path, w)
    out = tmp_path / "review.xlsx"
    cbs.links_review(db, lw.path, out, rules=w.rules)
    header, rows = _read_review(out)
    assert header[:8] == ["excel_row", "excel_company", "company_code", "debtor_code", "customer_name", "customer_id", "reason", "link"]
    mine = [r for r in rows if r["excel_row"] == 2 and r["company_code"] == "ZSRT"]
    assert 1 < len(mine) <= 5
    assert mine[0]["debtor_code"] == "ZZT-D-EXACT" and mine[1]["debtor_code"] == "ZZT-D-NEAR"
    assert all(r["excel_company"] == "ZZT Dealer A Sdn Bhd" and r["link"] in (None, "") and r["reason"] for r in mine)
    # only the row's own companies: the Mocha customer is not offered to a Sorento-only row
    assert "ZZT-M-1" not in {r["debtor_code"] for r in rows}
    assert mine[0]["customer_id"] == lw.exact.id


def test_u19_review_marks_an_already_linked_candidate(db, tmp_path, w):
    lw = _link_world(db, tmp_path, w)
    db.add(RespondContactCustomer(contact_id=lw.c1.id, customer_id=lw.exact.id, company_id=w.a.id, source="manual"))
    db.flush()
    out = tmp_path / "review.xlsx"
    cbs.links_review(db, lw.path, out, rules=w.rules)
    _, rows = _read_review(out)
    exact = next(r for r in rows if r["debtor_code"] == "ZZT-D-EXACT")
    assert "already linked" in exact["reason"]


def test_u19_review_offers_each_company_of_a_two_company_row(db, tmp_path, w):
    lw = _link_world(db, tmp_path, w)
    path = _book(tmp_path, [_row([("ZZT P1", raw(1), BO)], under="Sorento Sdn Bhd, Mocha Sdn Bhd")], name="two.xlsx")
    out = tmp_path / "review.xlsx"
    cbs.links_review(db, path, out, rules=w.rules)
    _, rows = _read_review(out)
    assert {r["company_code"] for r in rows} == {"ZSRT", "ZMCH"}


def test_u20_nothing_is_linked_by_review_or_by_run_apply(db, tmp_path, w):
    lw = _link_world(db, tmp_path, w)
    cbs.links_review(db, lw.path, tmp_path / "review.xlsx", rules=w.rules)
    _run(db, w, lw.path, apply=True, actor_id=w.actor)
    assert _links(db) == []


def test_u21_ticked_rows_link_to_every_found_contact(db, tmp_path, w):
    lw = _link_world(db, tmp_path, w)
    review, reviewed = tmp_path / "review.xlsx", tmp_path / "reviewed.xlsx"
    cbs.links_review(db, lw.path, review, rules=w.rules)
    _tick(review, reviewed, {"ZZT-D-EXACT"})
    res = cbs.apply_links(db, lw.path, reviewed, actor_id=w.actor, rules=w.rules)
    assert (res["linked"], res["skipped_existing"], res["failed"]) == (2, 0, [])
    links = _links(db)
    assert sorted(l.contact_id for l in links) == sorted([lw.c1.id, lw.c2.id])
    assert all(l.customer_id == lw.exact.id and l.source == "bulk_setup" and l.company_id == w.a.id and l.is_primary is False for l in links)
    assert len(_audits(db, event="contact.customer_linked_bulk")) == 2


@pytest.mark.parametrize("mark", ["Y", "YES", "x", "TRUE", "1", " yes "])
def test_u21_every_tick_spelling_counts(db, tmp_path, w, mark):
    lw = _link_world(db, tmp_path, w)
    review, reviewed = tmp_path / "review.xlsx", tmp_path / "reviewed.xlsx"
    cbs.links_review(db, lw.path, review, rules=w.rules)
    _tick(review, reviewed, {"ZZT-D-NEAR"}, mark=mark)
    assert cbs.apply_links(db, lw.path, reviewed, actor_id=w.actor, rules=w.rules)["linked"] == 2


@pytest.mark.parametrize("mark", ["N", "no", "", "maybe"])
def test_u21_unticked_rows_are_ignored(db, tmp_path, w, mark):
    lw = _link_world(db, tmp_path, w)
    review, reviewed = tmp_path / "review.xlsx", tmp_path / "reviewed.xlsx"
    cbs.links_review(db, lw.path, review, rules=w.rules)
    _tick(review, reviewed, {"ZZT-D-EXACT"}, mark=mark)
    assert cbs.apply_links(db, lw.path, reviewed, actor_id=w.actor, rules=w.rules)["linked"] == 0
    assert _links(db) == []


def test_u21_existing_pair_skipped_and_apply_is_idempotent(db, tmp_path, w):
    lw = _link_world(db, tmp_path, w)
    db.add(RespondContactCustomer(contact_id=lw.c1.id, customer_id=lw.exact.id, company_id=w.a.id, source="manual"))
    db.flush()
    review, reviewed = tmp_path / "review.xlsx", tmp_path / "reviewed.xlsx"
    cbs.links_review(db, lw.path, review, rules=w.rules)
    _tick(review, reviewed, {"ZZT-D-EXACT"})
    first = cbs.apply_links(db, lw.path, reviewed, actor_id=w.actor, rules=w.rules)
    assert (first["linked"], first["skipped_existing"]) == (1, 1)
    second = cbs.apply_links(db, lw.path, reviewed, actor_id=w.actor, rules=w.rules)
    assert (second["linked"], second["skipped_existing"]) == (0, 2)
    assert len(_links(db)) == 2
    assert len(_audits(db, event="contact.customer_linked_bulk")) == 1


def test_u22_vanished_customer_is_reported_others_still_link(db, tmp_path, w):
    lw = _link_world(db, tmp_path, w)
    review, reviewed = tmp_path / "review.xlsx", tmp_path / "reviewed.xlsx"
    cbs.links_review(db, lw.path, review, rules=w.rules)
    _tick(review, reviewed, {"ZZT-D-EXACT", "ZZT-D-NEAR"})
    db.delete(lw.near)
    db.flush()
    res = cbs.apply_links(db, lw.path, reviewed, actor_id=w.actor, rules=w.rules)
    assert res["linked"] == 2 and len(res["failed"]) == 1
    assert "ZZT-D-NEAR" in str(res["failed"][0])
    assert {l.customer_id for l in _links(db)} == {lw.exact.id}
